"""Document analyzer — parses project documents and extracts testing requirements via LLM.

Supports: PDF, DOCX, MD, TXT, HTML, RST.
"""

import io
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class AnalyzedDocument:
    """Result of document analysis."""

    doc_category: str = "other"  # requirement_spec, test_standard, code_style_guide, scoring_rubric
    key_requirements: list[str] = field(default_factory=list)
    test_criteria: list[str] = field(default_factory=list)
    scoring_rules: list[dict] = field(default_factory=list)
    code_style_rules: list[str] = field(default_factory=list)
    referenced_functions: list[str] = field(default_factory=list)
    edge_cases: list[str] = field(default_factory=list)
    environment_requirements: dict = field(default_factory=dict)
    uncertainty_notes: list[str] = field(default_factory=list)
    raw_text: str = ""


class DocumentAnalyzer:
    """Analyzes project documents to extract testing requirements.

    Uses LLM to understand structured and unstructured documents
    (PDF, DOCX, MD, TXT) and extracts key information for test planning.
    """

    SUPPORTED_FORMATS = {
        ".pdf": "PDF",
        ".docx": "Word",
        ".md": "Markdown",
        ".txt": "Text",
        ".html": "HTML",
        ".rst": "reStructuredText",
    }

    @classmethod
    def extract_text(cls, file_path: str) -> str:
        """Extract raw text from a document file.

        Args:
            file_path: Path to the document file.

        Returns:
            Extracted text content.

        Raises:
            ValueError: If the file format is unsupported.
            FileNotFoundError: If the file does not exist.
        """
        file_path = Path(file_path)
        ext = file_path.suffix.lower()

        if not file_path.exists():
            raise FileNotFoundError(f"Document not found: {file_path}")

        if ext == ".pdf":
            return cls._read_pdf(str(file_path))
        elif ext == ".docx":
            return cls._read_docx(str(file_path))
        elif ext in (".md", ".txt", ".html", ".rst"):
            return file_path.read_text(encoding="utf-8")
        else:
            raise ValueError(f"Unsupported document format: {ext}")

    @classmethod
    def _read_pdf(cls, file_path: str) -> str:
        """Read text from a PDF file using PyPDF2."""
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except ImportError:
            raise ImportError("PyPDF2 is required for PDF reading. Install: pip install PyPDF2")

    @classmethod
    def _read_docx(cls, file_path: str) -> str:
        """Read text from a DOCX file using python-docx."""
        try:
            from docx import Document
            doc = Document(file_path)
            text = "\n".join(para.text for para in doc.paragraphs)
            return text
        except ImportError:
            raise ImportError("python-docx is required. Install: pip install python-docx")

    @classmethod
    async def analyze_with_llm(
        cls,
        document_text: str,
        provider_manager=None,
        additional_context: str = "",
    ) -> AnalyzedDocument:
        """Analyze document content using LLM.

        Args:
            document_text: Extracted text from the document.
            provider_manager: ProviderManager instance for LLM calls.
            additional_context: Optional context (e.g., project description).

        Returns:
            AnalyzedDocument with extracted requirements.
        """
        from src.llm.base import BaseProvider

        context_info = ""
        if additional_context:
            context_info = f"\n\n## 补充说明\n{additional_context}"

        prompt = f"""你是一个专业的软件工程需求分析专家。请分析以下项目文档，提取代码检测所需的关键信息。

## 文档内容
{document_text}
{context_info}

请提取以下信息（只输出 JSON 格式）：

```json
{{
  "doc_category": "requirement_spec|test_standard|code_style_guide|scoring_rubric|other",
  "key_requirements": ["需求1", "需求2", ...],
  "test_criteria": ["测试标准1", ...],
  "scoring_rules": [
    {{"dimension": "功能正确性", "weight": 40, "description": "..."}},
    ...
  ],
  "code_style_rules": ["规范1", ...],
  "referenced_functions": ["函数签名1", ...],
  "edge_cases": ["边界条件1", ...],
  "environment_requirements": {{"language": "", "dependencies": [], "test_framework": ""}},
  "uncertainty_notes": ["文档中不明确的地方"]
}}
```
"""

        if provider_manager is None:
            from src.llm.provider import get_provider_manager
            provider_manager = get_provider_manager()

        try:
            response = await provider_manager.chat_with_fallback(
                messages=[BaseProvider.user(prompt)],
                temperature=0.3,
                max_tokens=4096,
            )

            content = response.content.strip()
            # Extract JSON from code blocks if present
            if "```json" in content:
                content = content.split("```json")[1].split("```")[0].strip()
            elif "```" in content:
                content = content.split("```")[1].split("```")[0].strip()

            data = json.loads(content)

            return AnalyzedDocument(
                doc_category=data.get("doc_category", "other"),
                key_requirements=data.get("key_requirements", []),
                test_criteria=data.get("test_criteria", []),
                scoring_rules=data.get("scoring_rules", []),
                code_style_rules=data.get("code_style_rules", []),
                referenced_functions=data.get("referenced_functions", []),
                edge_cases=data.get("edge_cases", []),
                environment_requirements=data.get("environment_requirements", {}),
                uncertainty_notes=data.get("uncertainty_notes", []),
                raw_text=document_text[:5000],
            )
        except (json.JSONDecodeError, KeyError) as e:
            logger.error(f"Failed to parse LLM analysis response: {e}")
            return AnalyzedDocument(raw_text=document_text[:5000])
        except Exception as e:
            logger.error(f"LLM analysis failed: {e}")
            return AnalyzedDocument(raw_text=document_text[:5000])